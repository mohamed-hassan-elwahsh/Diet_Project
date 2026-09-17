from fastapi import FastAPI, HTTPException, UploadFile, File, Form
from fastapi.responses import HTMLResponse, FileResponse
from pydantic import BaseModel
from typing import List, Optional
import os
import datetime
import shutil
import tempfile
import docx
import json
import re
import google.generativeai as genai
from diet_engine import load_default_plan, apply_adjustments
from docx_generator import generate_patient_diet_docx

app = FastAPI(title="Diet Assistant AI")

OUTPUT_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "output"))
os.makedirs(OUTPUT_DIR, exist_ok=True)
DATA_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "data"))
MASTER_PLAN_TXT = os.path.join(DATA_DIR, "master_meal_plan.txt")

class GenerateRequest(BaseModel):
    patient_name: str
    plan_date: str = ""
    weekly_plan: list = None

def extract_all_text(filepath):
    doc = docx.Document(filepath)
    texts = [node.text.strip() for node in doc.element.xpath('//w:t') if node.text.strip()]
    return " | ".join(texts)

def extract_plan_from_docx(filepath):
    try:
        doc = docx.Document(filepath)
        days_data = []
        day_names = ['السبت', 'الأحد', 'الإثنين', 'الثلاثاء', 'الأربعاء', 'الخميس', 'الجمعة']
        
        for t_idx, t in enumerate(doc.tables[:7]):
            day_info = {
                'day': day_names[t_idx] if t_idx < len(day_names) else f'Day {t_idx+1}',
                'meals': []
            }
            for r in t.rows[1:]:
                cells = [c.text.strip() for c in r.cells]
                meal_type = cells[1] if len(cells) > 1 else ''
                timing = cells[2] if len(cells) > 2 else ''
                diet_content = cells[3] if len(cells) > 3 else ''
                day_info['meals'].append({
                    'meal_type': meal_type,
                    'timing': timing,
                    'content': diet_content
                })
            days_data.append(day_info)
        
        if len(days_data) >= 7:
            return days_data
    except Exception as e:
        print("Error parsing old diet docx tables:", e)
    
    return load_default_plan()

def get_master_plan_text():
    if os.path.exists(MASTER_PLAN_TXT):
        with open(MASTER_PLAN_TXT, 'r', encoding='utf-8') as f:
            return f.read()
    return ""

@app.post("/api/analyze-patient")
async def analyze_patient(
    api_key: str = Form(None), 
    info_file: UploadFile = File(None),
    diet_file: UploadFile = File(None)  # Changed to accept single latest file from frontend
):
    try:
        if not api_key:
            return {"success": False, "error": "يرجى إدخال مفتاح Gemini API Key في أعلى الصفحة."}
        if not info_file:
            return {"success": False, "error": "يرجى اختيار ملف بيانات المريض أولاً."}

        genai.configure(api_key=api_key)
        temp_dir = tempfile.gettempdir()
        
        info_filename = os.path.basename(info_file.filename or "uploaded.docx")
        info_path = os.path.join(temp_dir, "info_" + info_filename)
        with open(info_path, "wb") as buffer:
            shutil.copyfileobj(info_file.file, buffer)
            
        old_diet_path = None
        
        if diet_file and diet_file.filename.endswith('.docx'):
            latest_diet_filename = os.path.basename(diet_file.filename)
            old_diet_path = os.path.join(temp_dir, "diet_" + latest_diet_filename)
            print("Processing latest diet file:", latest_diet_filename)
            with open(old_diet_path, "wb") as buffer:
                shutil.copyfileobj(diet_file.file, buffer)
            
        raw_text = extract_all_text(info_path)
        
        if old_diet_path:
            base_plan = extract_plan_from_docx(old_diet_path)
        else:
            base_plan = load_default_plan()
            
        formatted_base_plan = []
        for day in base_plan:
            formatted_day = {"day": day["day"], "meals": []}
            for m in day["meals"]:
                formatted_day["meals"].append({
                    "meal_type": m["meal_type"],
                    "timing": m["timing"],
                    "original_content": m.get("content", "")
                })
            formatted_base_plan.append(formatted_day)
            
        master_plan_db = get_master_plan_text()
        
        model = genai.GenerativeModel('gemini-3.6-flash')
        prompt = f"""
أنت مساعد ذكي ومحترف لدكتور تغذية.
إليك النص الخام المستخرج من ملف المريض (يحتوي على بياناته والـ Complaint أو تعليمات الدكتور):
"{raw_text}"

وإليك خطة الوجبات (JSON) للدايت القديم للمريض:
{json.dumps(formatted_base_plan, ensure_ascii=False)}

========================
[قاعدة بيانات الماستر بلان Master Meal Plan]
هذه هي الوجبات المتوفرة لديك للاختيار منها كبدائل مقسمة لجداول وصفوف:
{master_plan_db}
========================

التعليمات الهامة بدقة:
1. ابحث في نص المريض عن أحدث جلسة زمنياً واستخرج تاريخها، ثم استخرج "الشكوى / التعديلات" الخاصة بها.
2. استخرج "اسم المريض".
3. بالنسبة للوجبات: اقرأ `original_content`. المطلوب منك هو توفير **قائمة من 3 إلى 5 اقتراحات (بدائل مكافئة)** لكل وجبة.
4. يجب أن يتم استخراج هذه الاقتراحات حصراً من [قاعدة بيانات الماستر بلان] وتكون متوافقة مع جوهر الدايت والشكوى. (مثلاً إذا كانت الشكوى خالية من الكارب، اختر فقط البدائل الخالية من الكارب).
5. لكل اقتراح تقوم بوضعه، يجب أن تملأ حقل `content` (وهو محتوى الوجبة المقترحة)، وحقل `source` (مكانها الدقيق مثل: 'الماستر بلان: جدول رقم X، صف Y').

أخرج الرد بصيغة JSON فقط، مطابق لهذا الهيكل بالضبط:
{{
  "patient_name": "اسم المريض المستخرج",
  "last_session_date": "تاريخ أحدث جلسة مستخرج",
  "last_complaint": "نص شكوى الجلسة الأخيرة فقط",
  "suggested_plan": [
    {{
      "day": "السبت",
      "meals": [
        {{
          "meal_type": "...",
          "timing": "...",
          "original_content": "...",
          "suggestions": [
             {{ "content": "محتوى الاقتراح الأول", "source": "جدول رقم X، صف Y" }},
             {{ "content": "محتوى الاقتراح الثاني", "source": "جدول رقم X، صف Z" }}
          ]
        }}
      ]
    }}
  ]
}}
لا تكتب أي نص قبل أو بعد الـ JSON.
"""
        response = model.generate_content(prompt)
        txt = response.text.strip()
        
        if txt.startswith("```json"): txt = txt[7:]
        elif txt.startswith("```"): txt = txt[3:]
        if txt.endswith("```"): txt = txt[:-3]
        
        ai_data = json.loads(txt.strip())
        
        return {
            "success": True,
            "patient_name": ai_data.get("patient_name", ""),
            "last_session_date": ai_data.get("last_session_date", ""),
            "complaint": ai_data.get("last_complaint", ""),
            "suggested_plan": ai_data.get("suggested_plan", [])
        }
    except Exception as e:
        import traceback
        traceback.print_exc()
        return {"success": False, "error": str(e)}
    finally:
        try:
            if 'info_path' in locals() and os.path.exists(info_path): os.remove(info_path)
            if 'old_diet_path' in locals() and old_diet_path and os.path.exists(old_diet_path): os.remove(old_diet_path)
        except:
            pass

@app.get("/api/default-plan")
def get_default_plan():
    return []

@app.post("/api/generate-doc")
def generate_doc(req: GenerateRequest):
    if not req.patient_name.strip():
        raise HTTPException(status_code=400, detail="يرجى إدخال اسم المريض")
    
    plan_date = req.plan_date.strip() or datetime.date.today().strftime("%d-%m-%Y")
    
    plan_to_use = []
    for day in req.weekly_plan:
        new_day = {"day": day["day"], "meals": []}
        for m in day["meals"]:
            new_day["meals"].append({
                "meal_type": m["meal_type"],
                "timing": m["timing"],
                "content": m.get("selected_content", m.get("original_content", ""))
            })
        plan_to_use.append(new_day)
    
    if not plan_to_use:
        plan_to_use = load_default_plan()
    
    safe_name = "".join(c for c in req.patient_name if c.isalnum() or c in (' ', '-', '_')).strip()
    filename = f"{safe_name} - {plan_date}.docx"
    filepath = os.path.join(OUTPUT_DIR, filename)
    
    generate_patient_diet_docx(req.patient_name, plan_date, plan_to_use, filepath)
    
    return {
        "success": True,
        "filename": filename,
        "download_url": f"/download/{filename}",
        "full_path": filepath
    }

@app.get("/download/{filename}")
def download_file(filename: str):
    path = os.path.join(OUTPUT_DIR, filename)
    if os.path.exists(path):
        return FileResponse(path, filename=filename, media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    raise HTTPException(status_code=404, detail="File not found")

@app.get("/", response_class=HTMLResponse)
def index():
    html_path = os.path.join(os.path.dirname(__file__), "index.html")
    if os.path.exists(html_path):
        with open(html_path, "r", encoding="utf-8") as f:
            return f.read()
    return "<h1>Diet Assistant App</h1>"

if __name__ == "__main__":
    import uvicorn
    uvicorn.run(app, host="127.0.0.1", port=8000)
