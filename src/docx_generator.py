import docx
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import copy
import os
import re
import glob

TEMPLATE_PATH = os.path.join(os.path.dirname(__file__), "..", "templates", "weekly_diet_template.docx")
IMAGES_DIR = os.path.join(os.path.dirname(__file__), "..", "data", "images")

def find_image_for_source(source):
    if not source: return None
    # Parse numbers out of texts like "جدول رقم 3، صف 5"
    match = re.search(r'جدول.*?(\d+).*?صف.*?(\d+)', source)
    if match:
        table_idx = match.group(1)
        row_idx = match.group(2)
        # Search for any column image for this table and row
        pattern_jpg = os.path.join(IMAGES_DIR, f"table_{table_idx}_row_{row_idx}_col_*.jpg")
        pattern_png = os.path.join(IMAGES_DIR, f"table_{table_idx}_row_{row_idx}_col_*.png")
        
        images = glob.glob(pattern_jpg) + glob.glob(pattern_png)
        if images:
            return images[0] # Return the first found image (e.g. from col 1 or 2)
    return None

def set_cell_text_preserve_style(cell, text, source=""):
    """
    Sets text of a cell with clean, uniform formatting.
    Preserves images if no new image is provided.
    """
    new_image_path = find_image_for_source(source)
    
    drawings = []
    # If we don't have a new image from the master plan, save the old ones to preserve them
    if not new_image_path:
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                for child in run._r:
                    if child.tag.endswith('drawing') or child.tag.endswith('pict'):
                        drawings.append(copy.deepcopy(child))
                        
    # Clear the cell completely
    cell.text = ""
    
    # Write the new text with clean, standard formatting
    lines = text.split("\n")
    for i, line in enumerate(lines):
        if i == 0:
            p = cell.paragraphs[0]
        else:
            p = cell.add_paragraph()
            
        # Force right alignment for Arabic text to look professional
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        
        # Add BIDI (Right-to-Left) property to paragraph
        pPr = p._p.get_or_add_pPr()
        bidi = OxmlElement('w:bidi')
        bidi.set(qn('w:val'), '1')
        pPr.append(bidi)
        
        run = p.add_run(line.strip())
        run.font.name = "Segoe UI"
        run.font.size = Pt(13)
        run.font.bold = False
        
    # Inject images at the end of the cell
    if new_image_path:
        p = cell.paragraphs[-1]
        run = p.add_run('\n')
        try:
            # Insert the new extracted image
            run.add_picture(new_image_path, width=Inches(1.5))
        except Exception as e:
            print(f"Error inserting new image {new_image_path}: {e}")
            
    elif drawings:
        p = cell.paragraphs[-1]
        img_run = p.add_run('\n')
        for d in drawings:
            img_run._r.append(d)


def generate_patient_diet_docx(patient_name, plan_date, weekly_plan, output_path):
    doc = docx.Document(TEMPLATE_PATH)
    
    for day_idx, day_data in enumerate(weekly_plan[:7]):
        if day_idx < len(doc.tables):
            table = doc.tables[day_idx]
            for meal_idx, meal in enumerate(day_data.get("meals", [])):
                row_idx = meal_idx + 1
                if row_idx < len(table.rows):
                    row = table.rows[row_idx]
                    
                    if len(row.cells) > 1 and len(row.cells) > 2 and row.cells[1]._tc is row.cells[2]._tc:
                        # Cells are merged, combine text
                        combined_text = []
                        if meal.get("meal_type"): combined_text.append(meal["meal_type"])
                        if meal.get("timing"): combined_text.append(meal["timing"])
                        set_cell_text_preserve_style(row.cells[1], " - ".join(combined_text))
                    else:
                        # Not merged
                        if len(row.cells) > 1 and meal.get("meal_type"):
                            set_cell_text_preserve_style(row.cells[1], meal["meal_type"])
                        if len(row.cells) > 2 and meal.get("timing"):
                            set_cell_text_preserve_style(row.cells[2], meal["timing"])
                    
                    # Update content and pass the source to handle image swapping
                    if len(row.cells) > 3 and meal.get("content"):
                        source = meal.get("source", "")
                        set_cell_text_preserve_style(row.cells[3], meal["content"], source)

    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    doc.save(output_path)
    return output_path
