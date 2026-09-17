import docx
from docx.oxml.ns import qn
import os

print('Starting Image Extraction from Master Meal Plan...')
doc_path = 'Meal Plan.docx'
out_dir = 'data/images'
os.makedirs(out_dir, exist_ok=True)

doc = docx.Document(doc_path)

img_count = 0
for t_idx, table in enumerate(doc.tables):
    for r_idx, row in enumerate(table.rows):
        for c_idx, cell in enumerate(row.cells):
            for p in cell.paragraphs:
                for run in p.runs:
                    for child in run._r:
                        if child.tag.endswith('drawing') or child.tag.endswith('pict'):
                            blips = child.xpath('.//a:blip')
                            if blips:
                                embed = blips[0].get(qn('r:embed'))
                                if embed in doc.part.related_parts:
                                    image_part = doc.part.related_parts[embed]
                                    
                                    ext = 'jpg'
                                    if 'png' in image_part.content_type: ext = 'png'
                                    
                                    img_name = f"table_{t_idx+1}_row_{r_idx+1}_col_{c_idx+1}.{ext}"
                                    img_path = os.path.join(out_dir, img_name)
                                    
                                    # Save only if it doesn't exist
                                    if not os.path.exists(img_path):
                                        with open(img_path, 'wb') as f:
                                            f.write(image_part.blob)
                                        img_count += 1
                                        
print(f'Finished! Extracted {img_count} new images.')
